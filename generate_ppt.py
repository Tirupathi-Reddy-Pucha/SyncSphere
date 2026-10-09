import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    DARK_BG = RGBColor(15, 23, 42)
    PANEL_BG = RGBColor(30, 41, 59)
    CYAN_ACCENT = RGBColor(6, 182, 212)
    EMERALD_ACCENT = RGBColor(16, 185, 129)
    WHITE = RGBColor(255, 255, 255)
    LIGHT_GRAY = RGBColor(203, 213, 225)
    AMBER_ACCENT = RGBColor(245, 158, 11)

    def add_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = DARK_BG

    # Verified Direct Links
    AZURE_PORTAL_BASE = "https://portal.azure.com/#@cb.students.amrita.edu/resource"
    SUB = "a76241a6-5571-4c87-9829-a08e0ef360b9"
    RG = "rg-syncsphere-eastus2"
    APP_NAME = "app-syncsphere-api-2026"

    APP_SERVICE_PORTAL_LINK = f"{AZURE_PORTAL_BASE}/subscriptions/{SUB}/resourceGroups/{RG}/providers/Microsoft.Web/sites/{APP_NAME}/overview"
    RG_PORTAL_LINK = f"{AZURE_PORTAL_BASE}/subscriptions/{SUB}/resourceGroups/{RG}/overview"
    APP_CONFIG_PORTAL_LINK = f"{AZURE_PORTAL_BASE}/subscriptions/{SUB}/resourceGroups/{RG}/providers/Microsoft.Web/sites/{APP_NAME}/configuration"

    slides_data = [
        {
            "slide_type": "title",
            "title": "SyncSphere",
            "subtitle": "Cloud-Native Engineering Architecture & Proof of Work",
            "meta": "Live Azure Deployment • Infrastructure Verification • Source Code Proof Links",
            "live_url": "https://app-syncsphere-api-2026.azurewebsites.net/",
            "github_url": "https://github.com/Tirupathi-Reddy-Pucha/SyncSphere"
        },
        {
            "num": "1",
            "heading": "1. Frontend Architecture & React Single Page App",
            "tech": "React 18 • Vite • TailwindCSS • Lucide Icons",
            "host": "Azure App Service (React SPA Client Layer)",
            "description": "Production React frontend providing real-time collaboration boards, passwordless login interface, live Azure Blob Storage asset manager, and AI Security Scanner.",
            "proofs": [
                {"label": "🌐 Live Web Application URL", "url": "https://app-syncsphere-api-2026.azurewebsites.net/"},
                {"label": "🐙 Source Code (App.jsx on GitHub)", "url": "https://github.com/Tirupathi-Reddy-Pucha/SyncSphere/blob/main/frontend/src/App.jsx"},
                {"label": "☁️ Azure App Service Portal Location", "url": APP_SERVICE_PORTAL_LINK}
            ]
        },
        {
            "num": "2",
            "heading": "2. Backend REST API Microservices",
            "tech": "Node.js • Express.js • RESTful API Architecture",
            "host": "Azure App Service (app-syncsphere-api-2026)",
            "description": "High-performance Node.js Express backend serving REST API routes for project workspaces, live task board updates, blob uploads, and AI security diagnostics.",
            "proofs": [
                {"label": "⚡ Live Health API Endpoint URL", "url": "https://app-syncsphere-api-2026.azurewebsites.net/api/telemetry/health"},
                {"label": "🐙 Source Code (server.js on GitHub)", "url": "https://github.com/Tirupathi-Reddy-Pucha/SyncSphere/blob/main/backend/src/server.js"},
                {"label": "☁️ Azure App Service Configuration Portal", "url": APP_CONFIG_PORTAL_LINK}
            ]
        },
        {
            "num": "3",
            "heading": "3. Async Cloud Storage & Asset Manager",
            "tech": "@azure/storage-blob SDK • Buffer Streaming • MIME Auto-Detection",
            "host": "Azure Storage & Fallback Assets Container",
            "description": "Direct memory buffer streaming with automatic MIME preservation, direct preview download links, and instant blob deletion capabilities.",
            "proofs": [
                {"label": "🗄️ Live Asset Manager UI", "url": "https://app-syncsphere-api-2026.azurewebsites.net/"},
                {"label": "🐙 Source Code (azureBlobService.js on GitHub)", "url": "https://github.com/Tirupathi-Reddy-Pucha/SyncSphere/blob/main/backend/src/services/azureBlobService.js"},
                {"label": "☁️ Azure Resource Group Portal Location", "url": RG_PORTAL_LINK}
            ]
        },
        {
            "num": "4",
            "heading": "4. Security & Access Control (Passwordless Identity Lock)",
            "tech": "Passwordless Email Gateway • Client Session Lock • JWT Signatures",
            "host": "Session Security Module (frontend & backend middleware)",
            "description": "Enforces mandatory passwordless email login gateway. Permanently locks active session identity (clientSessionId) to prevent impersonation or name tampering.",
            "proofs": [
                {"label": "🔑 Live Passwordless Login Gateway UI", "url": "https://app-syncsphere-api-2026.azurewebsites.net/"},
                {"label": "🐙 Source Code (LoginScreen.jsx on GitHub)", "url": "https://github.com/Tirupathi-Reddy-Pucha/SyncSphere/blob/main/frontend/src/components/LoginScreen.jsx"},
                {"label": "🐙 Source Code (Header.jsx Locked Identity Badge)", "url": "https://github.com/Tirupathi-Reddy-Pucha/SyncSphere/blob/main/frontend/src/components/Header.jsx"}
            ]
        },
        {
            "num": "5",
            "heading": "5. Automated Azure AI Security Code Scanner",
            "tech": "Static Heuristic Analyzer • Regex Vulnerability Engine • 0-100 Scoring",
            "host": "Automated Upload Scanner & Interactive AI Code Auditor",
            "description": "Scans uploaded files & code snippets for hardcoded keys, HTTP connection strings, wildcard CORS, and public container access. Automatically outputs compliance scores.",
            "proofs": [
                {"label": "🛡️ Live AI Security Scanner UI", "url": "https://app-syncsphere-api-2026.azurewebsites.net/"},
                {"label": "🐙 Source Code (aiController.js on GitHub)", "url": "https://github.com/Tirupathi-Reddy-Pucha/SyncSphere/blob/main/backend/src/controllers/aiController.js"},
                {"label": "📁 Test Sample Files Directory on GitHub", "url": "https://github.com/Tirupathi-Reddy-Pucha/SyncSphere/tree/main/sample_security_test_files"}
            ]
        },
        {
            "num": "6",
            "heading": "6. Infrastructure Observability & Device-Aware Audit Stream",
            "tech": "Device Attribution (Desktop vs Mobile) • System Audit Log Stream",
            "host": "System Health & Audit Dashboard Component",
            "description": "Real-time telemetry monitoring for Azure services alongside a tamper-evident audit stream attributing all user actions to authenticated email & device hardware.",
            "proofs": [
                {"label": "📊 Live Telemetry & Audit Stream UI", "url": "https://app-syncsphere-api-2026.azurewebsites.net/"},
                {"label": "🐙 Source Code (TelemetryDashboard.jsx on GitHub)", "url": "https://github.com/Tirupathi-Reddy-Pucha/SyncSphere/blob/main/frontend/src/components/TelemetryDashboard.jsx"},
                {"label": "☁️ Azure Resource Group Portal Location", "url": RG_PORTAL_LINK}
            ]
        },
        {
            "num": "7",
            "heading": "7. Infrastructure as Code & Continuous Deployment Pipeline",
            "tech": "Azure Kudu Zip Deployment • Python CLI Automation • Git Version Control",
            "host": "GitHub Repository & Azure App Service Deployment API",
            "description": "Automated deployment pipeline creating zip deployment packages for React SPA build & Node.js backend, deploying cleanly to Azure App Service.",
            "proofs": [
                {"label": "🐙 Main GitHub Repository (SyncSphere)", "url": "https://github.com/Tirupathi-Reddy-Pucha/SyncSphere"},
                {"label": "🐍 Deployment Script (create_webapp.py on GitHub)", "url": "https://github.com/Tirupathi-Reddy-Pucha/SyncSphere/blob/main/create_webapp.py"},
                {"label": "📄 Implementation Plan PDF Report on GitHub", "url": "https://github.com/Tirupathi-Reddy-Pucha/SyncSphere/blob/main/Project_Collaboration_Workspace_Implementation_Plan.pdf"}
            ]
        }
    ]

    for item in slides_data:
        slide = prs.slides.add_slide(blank_layout)
        add_background(slide)

        if item.get("slide_type") == "title":
            tb = slide.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(11.333), Inches(4.5))
            tf = tb.text_frame
            tf.word_wrap = True

            p1 = tf.paragraphs[0]
            p1.text = item["title"]
            p1.font.size = Pt(48)
            p1.font.bold = True
            p1.font.color.rgb = CYAN_ACCENT
            p1.alignment = PP_ALIGN.CENTER

            p2 = tf.add_paragraph()
            p2.text = item["subtitle"]
            p2.font.size = Pt(24)
            p2.font.color.rgb = WHITE
            p2.alignment = PP_ALIGN.CENTER
            p2.space_before = Pt(15)

            p3 = tf.add_paragraph()
            p3.text = item["meta"]
            p3.font.size = Pt(16)
            p3.font.color.rgb = LIGHT_GRAY
            p3.alignment = PP_ALIGN.CENTER
            p3.space_before = Pt(20)

            btn_box = slide.shapes.add_textbox(Inches(2.0), Inches(5.2), Inches(9.333), Inches(1.5))
            btf = btn_box.text_frame
            btf.word_wrap = True

            bp1 = btf.paragraphs[0]
            bp1.alignment = PP_ALIGN.CENTER
            bp1.text = "🌐 Live App: " + item["live_url"]
            bp1.font.size = Pt(16)
            bp1.font.bold = True
            bp1.font.color.rgb = EMERALD_ACCENT

            bp2 = btf.add_paragraph()
            bp2.alignment = PP_ALIGN.CENTER
            bp2.text = "🐙 GitHub: " + item["github_url"]
            bp2.font.size = Pt(16)
            bp2.font.bold = True
            bp2.font.color.rgb = CYAN_ACCENT
            bp2.space_before = Pt(10)

        else:
            shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.6), Inches(11.733), Inches(6.3))
            shape.fill.solid()
            shape.fill.fore_color.rgb = PANEL_BG
            shape.line.color.rgb = CYAN_ACCENT
            shape.line.width = Pt(1.5)

            tb = slide.shapes.add_textbox(Inches(1.2), Inches(0.8), Inches(10.933), Inches(1.2))
            tf = tb.text_frame
            tf.word_wrap = True

            hp = tf.paragraphs[0]
            hp.text = item["heading"]
            hp.font.size = Pt(26)
            hp.font.bold = True
            hp.font.color.rgb = CYAN_ACCENT

            tp = tf.add_paragraph()
            tp.text = f"Tech Stack: {item['tech']}  |  Azure Host: {item['host']}"
            tp.font.size = Pt(13)
            tp.font.color.rgb = EMERALD_ACCENT
            tp.space_before = Pt(6)

            dp_box = slide.shapes.add_textbox(Inches(1.2), Inches(2.1), Inches(10.933), Inches(1.2))
            dtf = dp_box.text_frame
            dtf.word_wrap = True

            dp = dtf.paragraphs[0]
            dp.text = "Architecture & Feature Implementation:"
            dp.font.size = Pt(15)
            dp.font.bold = True
            dp.font.color.rgb = WHITE

            dp2 = dtf.add_paragraph()
            dp2.text = item["description"]
            dp2.font.size = Pt(14)
            dp2.font.color.rgb = LIGHT_GRAY
            dp2.space_before = Pt(4)

            pl_box = slide.shapes.add_textbox(Inches(1.2), Inches(3.6), Inches(10.933), Inches(3.0))
            pltf = pl_box.text_frame
            pltf.word_wrap = True

            lp0 = pltf.paragraphs[0]
            lp0.text = "🔗 Direct Location & Proof Links (Click to verify live):"
            lp0.font.size = Pt(16)
            lp0.font.bold = True
            lp0.font.color.rgb = AMBER_ACCENT

            for proof in item["proofs"]:
                lp = pltf.add_paragraph()
                lp.text = f"• {proof['label']}:"
                lp.font.size = Pt(13)
                lp.font.bold = True
                lp.font.color.rgb = WHITE
                lp.space_before = Pt(10)

                url_p = pltf.add_paragraph()
                url_p.text = f"  {proof['url']}"
                url_p.font.size = Pt(12)
                url_p.font.color.rgb = CYAN_ACCENT

    output_path = r"c:\Users\tirup\OneDrive\Desktop\cloud_project\SyncSphere_Architecture_Proof_Of_Work.pptx"
    prs.save(output_path)
    print(f"SUCCESS: PowerPoint saved to {output_path}")

if __name__ == "__main__":
    create_presentation()
